"""Deterministic synthetic CMP-P1 workbooks; never imports production inference."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).parent

DETAILS = (
    ("AWS", "svc-fictional-compute", "compute-test", "test-east-1", 2.50, 10, 25.00),
    ("AWS", "svc-fictional-storage", "object-test", "test-west-2", 1.25, 20, 25.00),
    ("AWS", "svc-fictional-compute", "compute-test", "test-west-2", 3.00, 5, 15.00),
)


def _save(
    name,
    headers,
    rows,
    *,
    title="ExampleCo-Test Synthetic Evidence",
    currency_note=None,
    footer=True,
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthetic Evidence"
    ws.append([title])
    ws.append(["Generated independently for CMP-P1 automated acceptance"])
    if currency_note:
        ws.append([currency_note])
    ws.append([])
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    if footer:
        ws.append([])
        ws.append(["Synthetic declared total", None, None, None, None, None, 65.00])
        ws.append(["Summary section - not detail"])
        ws.append(["svc-fictional-compute", 40.00])
        ws.append(["svc-fictional-storage", 25.00])
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    wb.save(ROOT / name)


def generate():
    ROOT.mkdir(parents=True, exist_ok=True)
    _save(
        "fixture_a_cloud_cost.xlsx",
        (
            "Cloud Provider",
            "Service",
            "Sub-Service",
            "Region",
            "Unit Price",
            "Quantity",
            "Extended Amount (USD)",
        ),
        DETAILS,
        currency_note="All monetary values: USD",
    )
    _save(
        "fixture_b_unfamiliar_headers.xlsx",
        (
            "Platform Mark",
            "Offering Label",
            "Flavor Band",
            "Geo Cell",
            "Alpha Rate",
            "Workload Units",
            "Line Extension (USD)",
        ),
        DETAILS,
        currency_note="Settlement denomination: USD",
    )
    _save(
        "fixture_c_ambiguous.xlsx",
        ("Category", "Metric A", "Metric B"),
        (("svc-fictional-a", 10, 20), ("svc-fictional-b", 5, 30)),
        footer=False,
    )
    _save(
        "fixture_d_mixed_currency.xlsx",
        (
            "Cloud Provider",
            "Service",
            "Region",
            "Unit Price",
            "Quantity",
            "Extended Amount",
            "Currency",
        ),
        (
            ("AWS", "svc-fictional-a", "test-east-1", 10, 1, 10, "USD"),
            ("AWS", "svc-fictional-b", "test-west-2", 10, 1, 10, "EUR"),
        ),
        footer=False,
    )
    _save(
        "fixture_e_inventory.xlsx",
        ("Asset Identifier", "Application", "Technology", "Owner", "Lifecycle"),
        (("resource-test-1", "app-fictional-shop", "db-fictional", "team-test", "ACTIVE"),),
        footer=False,
    )
    _save(
        "fixture_f_temporal_billing.xlsx",
        ("Cost Category", "Jan", "Feb", "Mar"),
        (
            ("svc-fictional-compute", 10.00, 12.00, 14.00),
            ("svc-fictional-storage", 5.00, 6.00, 7.00),
            ("Synthetic total", "=SUM(B5:B6)", "=SUM(C5:C6)", "=SUM(D5:D6)"),
        ),
        title="ExampleCo-Test Monthly billing amount",
        footer=False,
    )
    manifest = {
        "fixture_a_cloud_cost.xlsx": {
            "domain": "CLOUD_BILLING",
            "provider": "AWS",
            "measure": "EXTENDED_AMOUNT",
            "currency": "USD",
            "total": "65.00",
            "service": {"svc-fictional-compute": "40.00", "svc-fictional-storage": "25.00"},
            "region": {"test-east-1": "25.00", "test-west-2": "40.00"},
            "confirmation": ["domain", "measure"],
            "blocked": [],
        },
        "fixture_b_unfamiliar_headers.xlsx": {
            "domain": "CLOUD_BILLING",
            "provider": "AWS",
            "measure": "EXTENDED_AMOUNT",
            "currency": "USD",
            "total": "65.00",
            "service": {"svc-fictional-compute": "40.00", "svc-fictional-storage": "25.00"},
            "region": {"test-east-1": "25.00", "test-west-2": "40.00"},
            "confirmation": ["domain", "measure"],
            "blocked": [],
        },
        "fixture_c_ambiguous.xlsx": {
            "domain": "UNKNOWN",
            "provider": "UNKNOWN",
            "total": None,
            "confirmation": ["domain", "measure", "currency"],
            "blocked": ["domain", "measure", "currency"],
        },
        "fixture_d_mixed_currency.xlsx": {
            "domain": "CLOUD_BILLING",
            "provider": "AWS",
            "total": None,
            "confirmation": [],
            "blocked": ["mixed_currency"],
        },
        "fixture_e_inventory.xlsx": {
            "domain": "TECHNOLOGY_INVENTORY",
            "provider": "UNKNOWN",
            "total": None,
            "confirmation": [],
            "blocked": ["financial_activation"],
        },
        "fixture_f_temporal_billing.xlsx": {
            "domain": "CLOUD_BILLING",
            "provider": "UNKNOWN",
            "measure": "EXTENDED_AMOUNT",
            "currency": "HUMAN_CONFIRMATION_REQUIRED",
            "total": "54.00",
            "service": {
                "svc-fictional-compute": "36.00",
                "svc-fictional-storage": "18.00"
            },
            "region": {},
            "confirmation": ["domain", "measure", "currency"],
            "blocked": [],
        },
    }
    (ROOT / "ground_truth.json").write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    generate()
