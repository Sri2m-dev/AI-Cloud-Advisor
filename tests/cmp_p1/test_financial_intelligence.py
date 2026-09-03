from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID

import pytest
from openpyxl import load_workbook

from auth.authenticated_tenant import AuthenticatedTenantContext
from services.enterprise_spend_service import EnterpriseSpendService
from services.prospect_data_intake_service import ProspectIntakeError, scan_upload
from universal_evidence.domain import DomainClass, ProviderClass
from universal_evidence.financial import (
    CanonicalFinancialRepository,
    analyze_admitted_financial_evidence,
    analyze_financial_evidence,
)
from universal_evidence.financial.models import EvidenceTable, FinancialConcept
from universal_evidence.financial.publication import FinancialPublicationError
from universal_evidence.pilot.admission import admit_uploaded_evidence

ROOT = Path("tests/fixtures/cmp_p1")
ORG = str(UUID(int=101))


def _context(org=ORG):
    return AuthenticatedTenantContext(
        org,
        "ExampleCo-Test",
        "user-test",
        "test@example.invalid",
        "super_admin",
        frozenset({"financial:read"}),
        org,
    )


def _table(name):
    workbook = load_workbook(ROOT / name, read_only=True, data_only=False)
    sheet = workbook.active
    header_row = next(
        i
        for i, row in enumerate(sheet.iter_rows(values_only=True), 1)
        if sum(value is not None for value in row) >= 3
    )
    headers = tuple(
        str(value or "")
        for value in next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    )
    rows = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        row = tuple(row[: len(headers)])
        if not any(value is not None for value in row):
            break
        rows.append(row)
    return EvidenceTable(
        ORG,
        ORG,
        "prospect-test",
        "analysis-" + name[8],
        "source-test",
        name,
        "sheet-1",
        sheet.title,
        headers,
        tuple(rows),
    )


def test_fixture_generation_is_semantically_deterministic():
    manifest = json.loads((ROOT / "ground_truth.json").read_text(encoding="utf-8"))
    assert set(manifest) == {
        f"fixture_{letter}_{suffix}.xlsx"
        for letter, suffix in (
            ("a", "cloud_cost"),
            ("b", "unfamiliar_headers"),
            ("c", "ambiguous"),
            ("d", "mixed_currency"),
            ("e", "inventory"),
            ("f", "temporal_billing"),
        )
    }
    assert all(
        "ExampleCo-Test" in load_workbook(ROOT / name, read_only=True).active["A1"].value
        for name in manifest
    )


def test_production_admission_path_feeds_financial_intelligence():
    now = datetime.now(timezone.utc)
    tenant = SimpleNamespace(
        tenant_id="prospect-synthetic",
        audit_id="audit-synthetic",
        expires_at=(now + timedelta(days=30)).isoformat(),
    )
    admission = admit_uploaded_evidence(
        tenant,
        filename="unfamiliar-evidence.xlsx",
        content=(ROOT / "fixture_b_unfamiliar_headers.xlsx").read_bytes(),
        now=now,
        tenant_context=_context().fabric_context,
    )
    result = analyze_admitted_financial_evidence(
        admission, confirm_domain=True, confirm_measure=True
    )
    assert result.authorized
    assert sum(row.amount for row in result.observations) == Decimal("65.00")


@pytest.mark.parametrize("name", ("fixture_a_cloud_cost.xlsx", "fixture_b_unfamiliar_headers.xlsx"))
def test_unknown_headers_reach_production_cost_consumer(name, tmp_path):
    table = _table(name)
    analysis = analyze_financial_evidence(table, confirm_domain=True, confirm_measure=True)
    assert analysis.authorized
    assert analysis.domain.domain is DomainClass.CLOUD_BILLING
    assert analysis.domain.provider is ProviderClass.AWS
    amount = next(
        item for item in analysis.proposals if item.concept is FinancialConcept.EXTENDED_AMOUNT
    )
    assert "RECONCILIATION" in amount.signals
    assert analysis.currency.currencies == ("USD",)

    repository = CanonicalFinancialRepository(tmp_path / "financial.db")
    repository.publish(analysis.observations, _context())
    repository.publish(analysis.observations, _context())
    service = EnterpriseSpendService(repository, cache_ttl_seconds=0)
    posture = service.get_financial_posture(_context())
    assert posture.cloud_spend == Decimal("65.00")
    assert {row["service"]: row["amount"] for row in service.get_spend_by_service(_context())} == {
        "svc-fictional-compute": Decimal("40.00"),
        "svc-fictional-storage": Decimal("25.00"),
    }
    assert {row["region"]: row["amount"] for row in service.get_spend_by_region(_context())} == {
        "test-east-1": Decimal("25.00"),
        "test-west-2": Decimal("40.00"),
    }
    assert all(row["observation_ids"] for row in service.get_spend_by_service(_context()))


def test_ambiguous_financial_evidence_fails_closed():
    result = analyze_financial_evidence(_table("fixture_c_ambiguous.xlsx"))
    assert not result.authorized
    assert result.domain.domain is DomainClass.UNKNOWN
    assert not result.observations


def test_mixed_currency_fails_closed():
    result = analyze_financial_evidence(
        _table("fixture_d_mixed_currency.xlsx"), confirm_domain=True, confirm_measure=True
    )
    assert not result.authorized
    assert result.currency.conflicted
    assert not result.observations


def test_non_financial_inventory_does_not_activate():
    result = analyze_financial_evidence(
        _table("fixture_e_inventory.xlsx"),
        confirm_domain=True,
        confirm_measure=True,
        confirmed_currency="USD",
    )
    assert result.domain.domain is DomainClass.TECHNOLOGY_INVENTORY
    assert not result.authorized
    assert not result.observations


def test_provider_neutral_temporal_billing_requires_currency_confirmation():
    now = datetime.now(timezone.utc)
    tenant = SimpleNamespace(
        tenant_id="prospect-temporal",
        audit_id="audit-temporal",
        expires_at=(now + timedelta(days=30)).isoformat(),
    )
    admission = admit_uploaded_evidence(
        tenant,
        filename="provider-neutral-periods.xlsx",
        content=(ROOT / "fixture_f_temporal_billing.xlsx").read_bytes(),
        now=now,
        tenant_context=_context().fabric_context,
    )
    proposed = analyze_admitted_financial_evidence(
        admission, confirm_domain=True, confirm_measure=True
    )
    assert proposed.domain.domain is DomainClass.CLOUD_BILLING
    assert proposed.domain.provider is ProviderClass.UNKNOWN
    assert not proposed.authorized
    assert proposed.currency.currencies == ()
    assert len(proposed.proposals) == 3

    governed = analyze_admitted_financial_evidence(
        admission,
        confirm_domain=True,
        confirm_measure=True,
        confirmed_currency="USD",
    )
    assert governed.authorized
    assert len(governed.observations) == 6
    assert sum(item.amount for item in governed.observations) == Decimal("54.00")
    assert {dict(item.dimensions)["service"] for item in governed.observations} == {
        "svc-fictional-compute",
        "svc-fictional-storage",
    }


def test_cached_formula_amount_requires_reconciliation_and_governance():
    table = EvidenceTable(
        ORG,
        ORG,
        "prospect-formula",
        "analysis-formula",
        "source-formula",
        "formula.xlsx",
        "sheet-formula",
        "Synthetic Evidence",
        ("Offering", "Unit Rate", "Workload Units", "Calculated Extension (USD)"),
        (("svc-fictional-a", 2, 3, 6), ("svc-fictional-b", 5, 4, 20)),
        (("workbook_context", "synthetic cloud billing"),),
        ((0, 3), (1, 3)),
        ((0, 3), (1, 3)),
    )
    result = analyze_financial_evidence(table, confirm_domain=True, confirm_measure=True)
    assert result.authorized
    assert result.reconciliation and result.reconciliation.ratio == 1
    assert len(result.observations) == 2
    assert all(
        "cached-formula-value" in item.reconciliation_evidence
        for item in result.observations
    )


@pytest.mark.parametrize(
    ("headers", "rows"),
    (
        (("Asset ID", "CPU", "Memory"), ((10001, 8, 32), (10002, 16, 64))),
        (("Ticket", "Count", "Risk Score"), (("T-1", 7, 82), ("T-2", 3, 45))),
        (("Timestamp", "Latency", "Percent"), (("2026-01-01", 20, 99.9), ("2026-01-02", 18, 99.8))),
        (("Component", "Version", "Quantity"), (("runtime", "3.11", 10), ("agent", "2.0", 20))),
    ),
)
def test_stronger_numeric_inference_rejects_non_financial_tables(headers, rows):
    table = EvidenceTable(
        ORG, ORG, "prospect-negative", "analysis-negative", "source-negative",
        "negative.xlsx", "sheet-negative", "Technical Metrics", headers, rows,
    )
    result = analyze_financial_evidence(
        table, confirm_domain=True, confirm_measure=True, confirmed_currency="USD"
    )
    assert result.domain.domain is not DomainClass.CLOUD_BILLING
    assert not result.authorized
    assert not result.observations


def test_cross_tenant_kill_switch_purge_and_restart(tmp_path):
    analysis = analyze_financial_evidence(
        _table("fixture_a_cloud_cost.xlsx"), confirm_domain=True, confirm_measure=True
    )
    database = tmp_path / "financial.db"
    repository = CanonicalFinancialRepository(database)
    other = str(UUID(int=202))
    with pytest.raises(FinancialPublicationError, match="cross-tenant"):
        repository.publish(analysis.observations, _context(other))
    with pytest.raises(FinancialPublicationError, match="kill switch"):
        CanonicalFinancialRepository(database, kill_switch=True).publish(
            analysis.observations, _context()
        )
    repository.publish(analysis.observations, _context())
    restarted = CanonicalFinancialRepository(database)
    assert EnterpriseSpendService(restarted).get_financial_posture(_context()).cloud_spend
    assert restarted.purge_analysis(_context(), "analysis-a") == 3
    assert (
        EnterpriseSpendService(restarted, cache_ttl_seconds=0)
        .get_financial_posture(_context())
        .has_data
        is False
    )


def test_xlsx_binary_metadata_is_bounded_without_allowing_active_content():
    def package(member):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types />")
            archive.writestr(member, b"synthetic-test-metadata")
        return stream.getvalue()

    safe = package("xl/printerSettings/printerSettings1.bin")
    assert scan_upload("safe.xlsx", safe)["status"] == "PASS"
    with pytest.raises(ProspectIntakeError, match="active or embedded"):
        scan_upload("unsafe.xlsx", package("xl/vbaProject.bin"))
