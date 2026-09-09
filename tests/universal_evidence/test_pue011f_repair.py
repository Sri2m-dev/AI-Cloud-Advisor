from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook

from universal_evidence.containers import decompose_xlsx
from universal_evidence.contracts import EvidenceAnalysisContext
from universal_evidence.financial import analyze_decomposition_financial
from universal_evidence.semantics import (
    SemanticRegionType,
    classify_decomposition,
    group_evidence_sets,
)


def _pipeline(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    context = EvidenceAnalysisContext("analysis", "source", "prospect", "org", "tenant")
    decomposition = decompose_xlsx(
        context=context,
        content=stream.getvalue(),
        admitted_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
        source_reference="authorized:synthetic",
        filename="neutral.xlsx",
    )
    classifications = classify_decomposition(decomposition)
    sets = group_evidence_sets(decomposition.regions, classifications)
    documents = analyze_decomposition_financial(
        decomposition=decomposition,
        classifications=classifications,
        evidence_sets=sets,
        content=stream.getvalue(),
    )
    return decomposition, classifications, sets, documents


def test_real_decomposition_reaches_f_without_manual_financial_input():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invoice"
    sheet.append(["Invoice No.", "INV-1"])
    sheet.append(["Currency", "USD"])
    sheet.append([])
    sheet.append(["Description", "Qty", "Unit Price", "Amount"])
    sheet.append(["Service", 2, 5, 10])
    sheet.append([None, None, "Subtotal", 10])
    sheet.append([None, None, "Tax", 2])
    sheet.append([None, None, "Total Due", 12])
    _, _, sets, documents = _pipeline(workbook)
    assert len(sets) == 1 and len(documents) == 1
    result = documents[0].analysis
    assert result.eligible_detail_total == 10
    assert result.tax_total == 2 and result.gross_payable == 12
    assert result.currency.governed and result.currency.currencies == ("USD",)
    assert all(item.state.value == "RECONCILED" for item in result.reconciliations)


def test_access_only_leads_access_while_true_license_leads_license():
    workbook = Workbook()
    access = workbook.active
    access.title = "Access"
    access.append(["User", "Console Access", "Role", "Status"])
    access.append(["a@example.com", "Enabled", "Admin", "Assigned"])
    license_sheet = workbook.create_sheet("Licenses")
    license_sheet.append(["User", "License Tier", "Seat", "Assigned Status"])
    license_sheet.append(["b@example.com", "Business", "S-1", "Assigned"])
    decomposition, classifications, _, _ = _pipeline(workbook)
    by_locator = {
        r.raw_structure_reference: c for r, c in zip(decomposition.regions, classifications)
    }
    access_types = tuple(item.semantic_type for item in by_locator["Access!A1:D2"].candidates)
    license_types = tuple(item.semantic_type for item in by_locator["Licenses!A1:D2"].candidates)
    assert access_types[0] is SemanticRegionType.ACCESS_ENTITLEMENT
    assert license_types[0] is SemanticRegionType.LICENSE_ASSIGNMENT


def test_symbol_only_currency_remains_unresolved():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Description", "Qty", "Unit Price", "Amount"])
    sheet.append(["Service", 1, 10, 10])
    sheet.append([None, None, "Subtotal", 10])
    sheet.append([None, None, "Tax", 2])
    sheet.append([None, None, "Total Due", 12])
    for cell in sheet[2] + sheet[3] + sheet[4] + sheet[5]:
        cell.number_format = "$#,##0.00"
    *_, documents = _pipeline(workbook)
    assert not documents[0].analysis.currency.governed
    assert documents[0].analysis.currency.currencies == ()
