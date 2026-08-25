"""PUE-001 structural profiling acceptance and semantic-boundary tests."""

from __future__ import annotations

import io
from dataclasses import fields
from datetime import datetime, timedelta, timezone
from time import perf_counter

import pytest
from openpyxl import Workbook

from services.prospect_data_intake_service import normalize_upload
from universal_evidence.contracts import EvidenceAnalysisContext, EvidenceSource
from universal_evidence.profiling import (
    FileProfile,
    PrimitiveType,
    ProfilerConfig,
    ProfileStatus,
    StructuralRole,
    WarningCode,
    profile_evidence,
)


def source() -> EvidenceSource:
    now = datetime(2026, 8, 25, 10, 0, tzinfo=timezone.utc)
    context = EvidenceAnalysisContext(
        "analysis-pue-1", "source-pue-1", "prospect-pue-1"
    )
    return EvidenceSource(
        context, "authorized-upload:source-pue-1", now, now + timedelta(days=30)
    )


def warning_codes(profile: FileProfile) -> set[WarningCode]:
    codes = {warning.code for warning in profile.warnings}
    for sheet in profile.sheets:
        codes.update(warning.code for warning in sheet.warnings)
    return codes


def workbook_bytes() -> bytes:
    workbook = Workbook()
    clean = workbook.active
    clean.title = "Clean Table"
    clean.append(["Reference", "Value", "Group", "Observed Date"])
    clean.append(["ref-001", 10.5, "A", datetime(2026, 1, 1)])
    clean.append(["ref-002", 20.5, "A", datetime(2026, 1, 2)])
    clean.append(["ref-003", 30.5, "B", datetime(2026, 1, 3)])

    leading = workbook.create_sheet("Leading Rows")
    leading["A1"] = "Export title"
    leading.merge_cells("A1:C1")
    leading.append([])
    leading.append([])
    leading.append(["Code", "Code", "Amount"])
    leading.append(["x-1", "first", 1])
    leading.append(["x-2", "second", 2])

    mixed = workbook.create_sheet("Mixed Types")
    mixed.append(["Mixed", "Category"])
    mixed.append([1, "A"])
    mixed.append(["two", "A"])
    mixed.append([3.5, "B"])

    formulas = workbook.create_sheet("Formula Data")
    formulas.append(["Input", "Calculated"])
    formulas.append([2, "=A2*2"])
    formulas.append([3, "=A3*2"])
    formulas.sheet_state = "hidden"

    workbook.create_sheet("Empty")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def test_csv_file_profile_preserves_headers_and_statistics():
    content = (
        b"Reference,Value,Group,Optional\n"
        b"ref-001,10.5,A,\nref-002,20.5,A,x\nref-003,30.5,B,\n"
    )
    profile = profile_evidence(source=source(), filename="mixed.csv", content=content)
    assert profile.status is ProfileStatus.COMPLETE
    assert profile.container_format == "CSV"
    assert profile.physical_record_count == 4
    assert profile.logical_data_row_count == 3
    assert [column.column.original_header for column in profile.sheets[0].columns] == [
        "Reference",
        "Value",
        "Group",
        "Optional",
    ]
    optional = profile.sheets[0].columns[3].primitive_profile
    assert optional.null_count == 2
    assert optional.non_null_count == 1
    assert optional.coverage == pytest.approx(1 / 3)
    assert optional.distinct_count == 1


def test_csv_delimiter_and_exact_duplicate_rows_are_observed():
    profile = profile_evidence(
        source=source(), filename="data.csv", content=b"A;B\n1;x\n1;x\n2;y\n"
    )
    assert profile.delimiter == ";"
    assert profile.sheets[0].exact_duplicate_row_count == 1


def test_primitive_types_and_structural_roles_are_business_neutral():
    content = (
        b"Ref,Number,When,Category,Enabled\n"
        b"r-1,1.5,2026-01-01,A,true\n"
        b"r-2,2.5,2026-01-02,A,false\n"
        b"r-3,3.5,2026-01-03,B,true\n"
        b"r-4,4.5,2026-01-04,B,false\n"
    )
    columns = (
        profile_evidence(source=source(), filename="data.csv", content=content)
        .sheets[0]
        .columns
    )
    assert columns[0].dominant_primitive_type is PrimitiveType.STRING
    assert StructuralRole.IDENTIFIER_LIKE in columns[0].structural_roles
    assert StructuralRole.MEASURE_LIKE in columns[1].structural_roles
    assert StructuralRole.DATE_LIKE in columns[2].structural_roles
    assert StructuralRole.CATEGORICAL_LIKE in columns[3].structural_roles
    assert StructuralRole.BOOLEAN_LIKE in columns[4].structural_roles


def test_mixed_primitives_are_explicit():
    profile = profile_evidence(
        source=source(), filename="mixed.csv", content=b"Mixed\n1\ntwo\n3.5\n"
    )
    column = profile.sheets[0].columns[0]
    assert column.primitive_profile.mixed_type is True
    assert WarningCode.MIXED_PRIMITIVE_TYPES in {
        warning.code for warning in column.warnings
    }


def test_duplicate_and_unnamed_headers_are_warnings_without_rewriting_headers():
    profile = profile_evidence(
        source=source(), filename="headers.csv", content=b"A,A,\n1,2,3\n"
    )
    assert WarningCode.DUPLICATE_HEADERS in warning_codes(profile)
    assert WarningCode.UNNAMED_COLUMNS in warning_codes(profile)
    assert [item.column.original_header for item in profile.sheets[0].columns] == [
        "A",
        "A",
        "",
    ]


def test_repeated_headers_and_inconsistent_width_are_observed():
    profile = profile_evidence(
        source=source(), filename="rows.csv", content=b"A,B\n1,2\nA,B\n3\n"
    )
    assert WarningCode.REPEATED_HEADER_ROWS in warning_codes(profile)
    assert WarningCode.INCONSISTENT_ROW_WIDTH in warning_codes(profile)


def test_irregular_xlsx_profiles_sheets_independently():
    profile = profile_evidence(
        source=source(), filename="irregular.xlsx", content=workbook_bytes()
    )
    assert len(profile.sheets) == 5
    by_name = {sheet.sheet.original_name: sheet for sheet in profile.sheets}
    assert by_name["Clean Table"].sheet.header_row_candidate == 1
    assert by_name["Leading Rows"].sheet.header_row_candidate == 4
    assert by_name["Leading Rows"].sheet.empty_leading_rows == 0
    assert by_name["Mixed Types"].columns[0].primitive_profile.mixed_type is True
    assert by_name["Empty"].status is ProfileStatus.PARTIAL
    assert WarningCode.EMPTY_SHEET in {
        warning.code for warning in by_name["Empty"].warnings
    }


def test_xlsx_hidden_formula_merged_and_duplicate_header_observations():
    profile = profile_evidence(
        source=source(), filename="irregular.xlsx", content=workbook_bytes()
    )
    codes = warning_codes(profile)
    assert WarningCode.HIDDEN_SHEET in codes
    assert WarningCode.FORMULAS_PRESENT in codes
    assert WarningCode.MERGED_CELLS in codes
    assert WarningCode.DUPLICATE_HEADERS in codes
    formula_sheet = next(
        sheet for sheet in profile.sheets if sheet.sheet.original_name == "Formula Data"
    )
    assert formula_sheet.sheet.has_formulas is True
    assert formula_sheet.columns[1].primitive_type_distribution == (
        (PrimitiveType.FORMULA, 2),
    )


def test_multiple_regions_are_partial_not_forced_into_one_clean_table():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["A", "B"])
    sheet.append([1, 2])
    sheet.append([])
    sheet.append(["C", "D"])
    sheet.append([3, 4])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    profile = profile_evidence(
        source=source(), filename="regions.xlsx", content=stream.getvalue()
    )
    assert profile.status is ProfileStatus.PARTIAL
    assert WarningCode.MULTIPLE_TABLE_REGIONS in warning_codes(profile)


def test_malformed_csv_is_contained():
    profile = profile_evidence(
        source=source(), filename="broken.csv", content=b'A,B\n"unterminated,1\n'
    )
    assert profile.status in {ProfileStatus.PARTIAL, ProfileStatus.MALFORMED}
    assert WarningCode.MALFORMED_RECORDS in warning_codes(profile)


def test_corrupt_xlsx_is_contained():
    profile = profile_evidence(
        source=source(), filename="broken.xlsx", content=b"not-a-workbook"
    )
    assert profile.status is ProfileStatus.MALFORMED
    assert WarningCode.MALFORMED_RECORDS in warning_codes(profile)


def test_sample_count_and_length_are_bounded_and_provenanced():
    config = ProfilerConfig(max_sample_values=2, max_string_sample_length=8)
    content = b"Value\nabcdefghijk\nsecondvalue\nthirdvalue\n"
    column = (
        profile_evidence(
            source=source(), filename="samples.csv", content=content, config=config
        )
        .sheets[0]
        .columns[0]
    )
    assert len(column.samples) == 2
    assert all(len(sample.display_value) <= 9 for sample in column.samples)
    assert column.samples[0].truncated is True
    assert column.samples[0].row_reference.file_id == column.column.file_id
    assert column.samples[0].row_reference.sheet_id == column.column.sheet_id
    assert column.samples[0].row_reference.context.analysis_id == "analysis-pue-1"


def test_high_entropy_sample_is_suppressed_without_pii_classification():
    token = b"A9f3K2m8Q1z7W4x6V0b5N2c9L8p7R4t1"
    profile = profile_evidence(
        source=source(), filename="safe.csv", content=b"Value\n" + token + b"\n"
    )
    column = profile.sheets[0].columns[0]
    assert column.samples == ()
    assert column.primitive_profile.sensitive_looking is True
    assert WarningCode.SAMPLE_SUPPRESSED in {
        warning.code for warning in column.warnings
    }


def test_resource_limits_return_partial_profiles():
    profile = profile_evidence(
        source=source(),
        filename="limited.csv",
        content=b"A,B\n1,2\n3,4\n5,6\n",
        config=ProfilerConfig(max_rows_profiled=1),
    )
    assert profile.status is ProfileStatus.PARTIAL
    assert WarningCode.RESOURCE_LIMIT_REACHED in warning_codes(profile)


def test_unsupported_format_is_explicit_and_does_not_search_paths():
    profile = profile_evidence(source=source(), filename="evidence.json", content=b"{}")
    assert profile.status is ProfileStatus.UNSUPPORTED_FORMAT
    assert profile.sheets == ()


def test_identical_input_produces_identical_structural_fingerprint():
    content = b"A,B\n1,x\n2,y\n"
    first = profile_evidence(source=source(), filename="same.csv", content=content)
    second = profile_evidence(source=source(), filename="same.csv", content=content)
    assert first.structural_fingerprint == second.structural_fingerprint


@pytest.mark.parametrize(
    ("filename", "header", "value"),
    [
        ("aws_cur_january.csv", "AWS Service", "EC2"),
        ("neutral.csv", "EC2 Cost", "10"),
        ("neutral.csv", "Application Owner", "Team A"),
        ("neutral.csv", "Resource", "i-0123456789abcdef"),
    ],
)
def test_business_names_and_values_never_emit_semantic_assignments(
    filename, header, value
):
    profile = profile_evidence(
        source=source(),
        filename=filename,
        content=f"{header}\n{value}\nother\nthird\n".encode(),
    )
    assert all("semantic" not in field.name for field in fields(type(profile)))
    assert all(
        "semantic" not in field.name
        for sheet in profile.sheets
        for column in sheet.columns
        for field in fields(type(column))
    )
    allowed_roles = set(StructuralRole)
    assert all(
        set(column.structural_roles) <= allowed_roles
        for sheet in profile.sheets
        for column in sheet.columns
    )


def test_profile_has_no_business_aggregate_surface():
    profile = profile_evidence(
        source=source(), filename="EC2 Cost.csv", content=b"EC2 Cost\n10\n20\n"
    )
    names = {field.name for field in fields(type(profile))}
    assert not names.intersection(
        {"total_spend", "cost_by_service", "currency", "provider"}
    )


def test_184_row_profiling_does_not_change_authoritative_cost_or_currency_behavior():
    costs = ["1"] * 183 + [str(861_828 - 183)]
    content = (
        "provider,service,cost\n" + "".join(f"AWS,EC2,{cost}\n" for cost in costs)
    ).encode()
    before = normalize_upload(
        "AWS billing/CUR-derived CSV", "CUR Jan 2026.csv", content
    )
    profile = profile_evidence(
        source=source(), filename="CUR Jan 2026.csv", content=content
    )
    after = normalize_upload("AWS billing/CUR-derived CSV", "CUR Jan 2026.csv", content)
    assert profile.logical_data_row_count == 184
    assert profile.sheets[0].sheet.row_count == 184
    assert not hasattr(profile, "total_spend")
    assert float(before["cost"].sum()) == float(after["cost"].sum()) == 861_828
    assert before["currency"].isna().all() and after["currency"].isna().all()


def test_representative_profile_has_bounded_runtime():
    rows = ["A,B,C"] + [f"id-{index},{index},{index % 5}" for index in range(2_000)]
    content = ("\n".join(rows) + "\n").encode()
    started = perf_counter()
    profile = profile_evidence(
        source=source(), filename="performance.csv", content=content
    )
    duration = perf_counter() - started
    assert profile.logical_data_row_count == 2_000
    assert duration < 5.0
