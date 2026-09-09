from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from universal_evidence.containers import XlsxDecompositionConfig, decompose_xlsx
from universal_evidence.contracts import EvidenceAnalysisContext
from universal_evidence.documents import (
    SecurityFinding,
    StructuralState,
    StructuralType,
)

NOW = datetime(2026, 1, 1, tzinfo=timezone.utc)


def context(tenant="tenant-a"):
    return EvidenceAnalysisContext("analysis", "source", "prospect", "org", tenant)


def save(workbook):
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def decompose(content, *, tenant="tenant-a", filename="neutral.xlsx", config=None):
    return decompose_xlsx(
        context=context(tenant),
        content=content,
        admitted_at=NOW,
        source_reference="authorized:synthetic",
        filename=filename,
        config=config,
    )


def add_table(sheet, row, column=1, labels=("Field A", "Field B", "Field C")):
    for offset, value in enumerate(labels):
        sheet.cell(row, column + offset, value)
    for row_offset in range(1, 4):
        sheet.cell(row + row_offset, column, f"r-{row_offset}")
        sheet.cell(row + row_offset, column + 1, row_offset)
        sheet.cell(row + row_offset, column + 2, row_offset * 2)


@pytest.mark.parametrize("start", [1, 9, 100])
def test_table_header_is_discovered_per_region_at_arbitrary_bounded_row(start):
    workbook = Workbook()
    add_table(workbook.active, start)
    result = decompose(save(workbook))
    assert len(result.regions) == 1
    assert result.regions[0].structural_type is StructuralType.TABULAR
    assert result.region_profiles[0][1].header_row == start


def test_multiple_vertical_and_side_by_side_tables_are_independent():
    workbook = Workbook()
    sheet = workbook.active
    add_table(sheet, 2, 1)
    add_table(sheet, 10, 1, ("One", "Two", "Three"))
    add_table(sheet, 2, 6, ("Left", "Middle", "Right"))
    result = decompose(save(workbook))
    bounds = {
        (
            item.location.start_row,
            item.location.end_row,
            item.location.start_column,
            item.location.end_column,
        )
        for item in result.regions
    }
    assert bounds == {(2, 5, 1, 3), (2, 5, 6, 8), (10, 13, 1, 3)}
    assert len({item.region_id for item in result.regions}) == 3


def test_metadata_table_and_summary_blocks_remain_separate():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Label A", "Value A"])
    sheet.append(["Label B", "Value B"])
    add_table(sheet, 5)
    sheet.cell(11, 1, "Measure A")
    sheet.cell(11, 2, 10)
    sheet.cell(12, 1, "Measure B")
    sheet.cell(12, 2, 20)
    result = decompose(save(workbook))
    assert [item.structural_type for item in result.regions] == [
        StructuralType.KEY_VALUE,
        StructuralType.TABULAR,
        StructuralType.SUMMARY,
    ]


def test_multiple_sheets_and_ambiguous_sparse_region_are_preserved():
    workbook = Workbook()
    add_table(workbook.active, 1)
    second = workbook.create_sheet("Arbitrary")
    second["D40"] = 7
    result = decompose(save(workbook))
    assert len(result.sheets) == 2
    assert result.regions[-1].structural_type is StructuralType.UNKNOWN
    assert result.regions[-1].status is StructuralState.UNRESOLVED


def test_formats_formulas_merges_hidden_content_and_named_ranges_are_observed():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Neutral Grid"
    add_table(sheet, 4)
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "Bounded heading"
    sheet["B5"].number_format = "$#,##0.00"
    sheet["C5"] = "=B5*2"
    sheet["A6"].number_format = "yyyy-mm-dd"
    sheet.row_dimensions[6].hidden = True
    sheet.column_dimensions["C"].hidden = True
    hidden = workbook.create_sheet("Background")
    hidden["A1"] = "Observation"
    hidden.sheet_state = "hidden"
    workbook.defined_names.add(DefinedName("BoundedArea", attr_text="'Neutral Grid'!$A$4:$C$7"))
    result = decompose(save(workbook))
    main = result.sheets[0]
    profile = next(
        profile
        for _, profile in result.region_profiles
        if "$#,##0.00" in profile.number_formats
    )
    assert "$#,##0.00" in profile.number_formats
    assert profile.formulas[0].cell_reference == "C5"
    assert profile.formulas[0].expression == "=B5*2"
    assert not profile.formulas[0].cached_value_available
    assert "A1:C1" in main.merged_ranges
    assert main.hidden_rows == (6,)
    assert main.hidden_columns == ("C",)
    assert main.named_ranges[0].name == "BoundedArea"
    assert any(
        item.finding is SecurityFinding.HIDDEN_CONTENT_PRESENT
        for item in result.container.security_diagnostics
    )


def test_identity_is_replay_stable_filename_independent_and_tenant_scoped():
    workbook = Workbook()
    add_table(workbook.active, 9)
    content = save(workbook)
    first = decompose(content, filename="first.xlsx")
    renamed = decompose(content, filename="renamed.xlsx")
    foreign = decompose(content, tenant="tenant-b")
    assert first.container.container_id == renamed.container.container_id
    assert [item.region_id for item in first.regions] == [
        item.region_id for item in renamed.regions
    ]
    assert first.container.container_id != foreign.container.container_id
    assert [item.region_id for item in first.regions] != [
        item.region_id for item in foreign.regions
    ]


def test_layout_and_label_mutation_does_not_require_semantic_dispatch():
    first_workbook = Workbook()
    add_table(first_workbook.active, 3, labels=("Alpha", "Beta", "Gamma"))
    second_workbook = Workbook()
    second_workbook.active.title = "Changed"
    add_table(second_workbook.active, 30, labels=("North", "East", "West"))
    first = decompose(save(first_workbook))
    second = decompose(save(second_workbook))
    assert first.regions[0].structural_type is StructuralType.TABULAR
    assert second.regions[0].structural_type is StructuralType.TABULAR
    assert not hasattr(first, "financial_observations")
    assert not hasattr(first, "source_facts")


def test_resource_exhaustion_is_partial_and_malformed_content_is_quarantined():
    workbook = Workbook()
    add_table(workbook.active, 1)
    limited = decompose(
        save(workbook), config=XlsxDecompositionConfig(max_cells_inspected=2)
    )
    malformed = decompose(b"not-a-workbook")
    assert limited.status is StructuralState.PARTIAL
    assert limited.container.resource_diagnostics.limit_exceeded
    assert malformed.status is StructuralState.QUARANTINED
    assert not malformed.container.publishable
    assert any(
        item.finding is SecurityFinding.MALFORMED
        for item in malformed.container.security_diagnostics
    )


def test_active_content_is_detected_and_quarantined_without_execution():
    workbook = Workbook()
    add_table(workbook.active, 1)
    source = zipfile.ZipFile(io.BytesIO(save(workbook)))
    stream = io.BytesIO()
    with source, zipfile.ZipFile(stream, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item.filename))
        target.writestr("xl/vbaProject.bin", b"synthetic-active-content")
    result = decompose(stream.getvalue())
    assert result.status is StructuralState.QUARANTINED
    assert not result.regions
    assert any(
        item.finding is SecurityFinding.MACRO_PRESENT
        for item in result.container.security_diagnostics
    )


def test_structural_output_never_emits_semantic_region_types():
    workbook = Workbook()
    add_table(workbook.active, 1, labels=("Invoice", "License", "Owner"))
    result = decompose(save(workbook))
    assert all(item.structural_type in set(StructuralType) for item in result.regions)
    prohibited = {"INVOICE", "LICENSE_ASSIGNMENT", "USAGE", "CLOUD_COST", "OWNER"}
    assert prohibited.isdisjoint({item.structural_type.value for item in result.regions})
