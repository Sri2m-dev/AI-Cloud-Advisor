"""Independent-sheet XLSX structural profiler using existing openpyxl dependency."""

from __future__ import annotations

import io
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from universal_evidence.contracts import EvidenceFile, EvidenceSheet, EvidenceSource
from universal_evidence.profiling.common import (
    build_column_profile,
    detect_header,
    duplicate_header_ordinals,
    exact_duplicate_count,
    nonempty,
    row_regions,
)
from universal_evidence.profiling.models import (
    FileProfile,
    ProfilerConfig,
    ProfileStatus,
    SheetProfile,
    StructuralWarning,
    WarningCode,
)


def _malformed(
    source: EvidenceSource, evidence_file: EvidenceFile, version: str, detail: str
) -> FileProfile:
    warning = StructuralWarning(WarningCode.MALFORMED_RECORDS, "file", detail)
    return FileProfile(
        evidence_file,
        ProfileStatus.MALFORMED,
        "XLSX",
        version,
        source.received_at,
        (),
        warnings=(warning,),
    )


def profile_xlsx(
    source: EvidenceSource,
    evidence_file: EvidenceFile,
    content: bytes,
    config: ProfilerConfig,
    profiler_version: str,
) -> FileProfile:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            expanded = sum(item.file_size for item in archive.infolist())
            if expanded > config.max_expanded_xlsx_size:
                warning = StructuralWarning(
                    WarningCode.RESOURCE_LIMIT_REACHED,
                    "file",
                    "expanded workbook size limit reached",
                )
                return FileProfile(
                    evidence_file,
                    ProfileStatus.PARTIAL,
                    "XLSX",
                    profiler_version,
                    source.received_at,
                    (),
                    warnings=(warning,),
                )
    except zipfile.BadZipFile:
        return _malformed(
            source, evidence_file, profiler_version, "invalid XLSX container"
        )
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile):
        return _malformed(
            source,
            evidence_file,
            profiler_version,
            "workbook could not be structurally opened",
        )

    file_warnings: list[StructuralWarning] = []
    workbook_limited = len(workbook.worksheets) > config.max_sheets
    worksheets = workbook.worksheets[: config.max_sheets]
    if workbook_limited:
        file_warnings.append(
            StructuralWarning(
                WarningCode.RESOURCE_LIMIT_REACHED,
                "file",
                "sheet profiling limit reached",
            )
        )
    profiles: list[SheetProfile] = []
    for sheet_ordinal, worksheet in enumerate(worksheets, start=1):
        sheet_id = f"{evidence_file.file_id}:sheet:{sheet_ordinal}"
        sheet_warnings: list[StructuralWarning] = []
        hidden = worksheet.sheet_state != "visible"
        if hidden:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.HIDDEN_SHEET, sheet_id, "sheet is not visible"
                )
            )
        if worksheet.merged_cells.ranges:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.MERGED_CELLS,
                    sheet_id,
                    f"{len(worksheet.merged_cells.ranges)} merged region(s) observed",
                )
            )
        max_row = min(
            worksheet.max_row or 0, config.max_rows_profiled + config.header_scan_rows
        )
        max_column = min(worksheet.max_column or 0, config.max_columns)
        limited = (worksheet.max_row or 0) > max_row or (
            worksheet.max_column or 0
        ) > max_column
        if limited:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.RESOURCE_LIMIT_REACHED,
                    sheet_id,
                    "sheet row or column profiling limit reached",
                )
            )
        rows: list[list[Any]] = []
        formula_matrix: list[list[bool]] = []
        for cells in worksheet.iter_rows(
            min_row=1, max_row=max_row, min_col=1, max_col=max_column
        ):
            rows.append([cell.value for cell in cells])
            formula_matrix.append([cell.data_type == "f" for cell in cells])
        populated_indexes = [
            index
            for index, row in enumerate(rows)
            if any(nonempty(value) for value in row)
        ]
        if not populated_indexes:
            warning = StructuralWarning(
                WarningCode.EMPTY_SHEET, sheet_id, "sheet contains no observable values"
            )
            sheet_warnings.append(warning)
            sheet = EvidenceSheet(
                source.context,
                evidence_file.file_id,
                sheet_id,
                worksheet.title,
                sheet_ordinal,
                hidden,
                0,
                0,
                warnings=tuple(item.code.value for item in sheet_warnings),
            )
            profiles.append(
                SheetProfile(
                    sheet,
                    ProfileStatus.PARTIAL,
                    None,
                    0.0,
                    "NO_NONEMPTY_ROWS",
                    (),
                    0,
                    tuple(sheet_warnings),
                )
            )
            continue
        first_populated, last_populated = populated_indexes[0], populated_indexes[-1]
        empty_leading = first_populated
        empty_trailing = len(rows) - last_populated - 1
        trimmed_rows = rows[: last_populated + 1]
        trimmed_formulas = formula_matrix[: last_populated + 1]
        header_row, header_confidence, header_method = detect_header(
            trimmed_rows, config
        )
        if header_row is None or header_confidence < 0.45:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.HEADER_UNCERTAIN,
                    sheet_id,
                    "header candidate has low structural confidence",
                )
            )
        candidate_index = (header_row - 1) if header_row else first_populated
        headers = trimmed_rows[candidate_index][:max_column]
        width = max((len(row) for row in trimmed_rows), default=0)
        headers = list(headers) + [None] * (width - len(headers))
        if duplicate_header_ordinals(headers):
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.DUPLICATE_HEADERS,
                    sheet_id,
                    "duplicate header labels observed",
                )
            )
        if any(not nonempty(value) for value in headers):
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.UNNAMED_COLUMNS,
                    sheet_id,
                    "one or more header cells are blank",
                )
            )
        formula_count = sum(sum(row) for row in trimmed_formulas)
        if formula_count:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.FORMULAS_PRESENT,
                    sheet_id,
                    f"{formula_count} formula cell(s) observed; formulas were not evaluated",
                )
            )
        regions = row_regions(trimmed_rows)
        if len(regions) > 1:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.MULTIPLE_TABLE_REGIONS,
                    sheet_id,
                    "multiple non-empty row regions separated by blanks",
                )
            )
        data_rows: list[tuple[int, list[Any], list[bool]]] = []
        repeated = 0
        header_signature = tuple(str(value).strip() for value in headers)
        for index in range(candidate_index + 1, len(trimmed_rows)):
            row = list(trimmed_rows[index]) + [None] * (
                width - len(trimmed_rows[index])
            )
            formulas = list(trimmed_formulas[index]) + [False] * (
                width - len(trimmed_formulas[index])
            )
            if not any(nonempty(value) for value in row):
                continue
            if tuple(str(value).strip() for value in row) == header_signature:
                repeated += 1
                continue
            data_rows.append((index + 1, row, formulas))
        if len(data_rows) > config.max_rows_profiled:
            data_rows = data_rows[: config.max_rows_profiled]
            limited = True
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.RESOURCE_LIMIT_REACHED,
                    sheet_id,
                    "logical data row profiling limit reached",
                )
            )
        if repeated:
            sheet_warnings.append(
                StructuralWarning(
                    WarningCode.REPEATED_HEADER_ROWS,
                    sheet_id,
                    f"{repeated} repeated header-like row(s) observed",
                )
            )
        columns = []
        for ordinal in range(width):
            column_id = f"{sheet_id}:column:{ordinal + 1}"
            values = [
                (row_number, row[ordinal], formulas[ordinal])
                for row_number, row, formulas in data_rows
            ]
            columns.append(
                build_column_profile(
                    context=source.context,
                    file_id=evidence_file.file_id,
                    sheet_id=sheet_id,
                    column_id=column_id,
                    original_header=headers[ordinal],
                    ordinal=ordinal + 1,
                    values=values,
                    config=config,
                )
            )
        all_warnings = tuple(
            sheet_warnings
            + [warning for column in columns for warning in column.warnings]
        )
        partial = limited or len(regions) > 1 or header_confidence < 0.45
        status = ProfileStatus.PARTIAL if partial else ProfileStatus.COMPLETE
        sheet = EvidenceSheet(
            source.context,
            evidence_file.file_id,
            sheet_id,
            worksheet.title,
            sheet_ordinal,
            hidden,
            len(data_rows),
            width,
            header_row,
            empty_leading,
            empty_trailing,
            bool(worksheet.merged_cells.ranges),
            bool(formula_count),
            tuple(warning.code.value for warning in all_warnings),
        )
        profiles.append(
            SheetProfile(
                sheet,
                status,
                (first_populated + 1, 1, last_populated + 1, width),
                header_confidence,
                header_method,
                tuple(columns),
                exact_duplicate_count([row for _, row, _ in data_rows]),
                all_warnings,
            )
        )
    workbook.close()
    statuses = {profile.status for profile in profiles}
    overall = (
        ProfileStatus.PARTIAL
        if workbook_limited or ProfileStatus.PARTIAL in statuses
        else ProfileStatus.COMPLETE
    )
    return FileProfile(
        evidence_file,
        overall,
        "XLSX",
        profiler_version,
        source.received_at,
        tuple(profiles),
        warnings=tuple(file_warnings),
    )
