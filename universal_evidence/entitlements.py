"""Shadow normalization for access, license assignment, and usage evidence."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from decimal import Decimal
from enum import Enum

from openpyxl import load_workbook

from universal_evidence.documents import (
    SpreadsheetLocation,
    structural_fingerprint,
)
from universal_evidence.semantics import SemanticRegionType

ENTITLEMENT_VERSION = "pue-011h.entitlement.v1"


class EntitlementConcept(str, Enum):
    ACCESS_ENTITLEMENT = "ACCESS_ENTITLEMENT"
    LICENSE_ENTITLEMENT = "LICENSE_ENTITLEMENT"
    LICENSE_ASSIGNMENT = "LICENSE_ASSIGNMENT"
    PURCHASED_QUANTITY = "PURCHASED_QUANTITY"
    ASSIGNED_QUANTITY = "ASSIGNED_QUANTITY"
    UNASSIGNED_QUANTITY = "UNASSIGNED_QUANTITY"
    ACTIVE_ASSIGNMENT = "ACTIVE_ASSIGNMENT"
    INACTIVE_ASSIGNMENT = "INACTIVE_ASSIGNMENT"
    USAGE_OBSERVATION = "USAGE_OBSERVATION"
    LAST_ACTIVITY = "LAST_ACTIVITY"
    USAGE_PERIOD = "USAGE_PERIOD"
    LICENSE_TIER = "LICENSE_TIER"
    LICENSE_PRICE = "LICENSE_PRICE"
    ENTITLEMENT_SCOPE = "ENTITLEMENT_SCOPE"
    ACCOUNT_IDENTIFIER = "ACCOUNT_IDENTIFIER"


class QuantityReconciliationState(str, Enum):
    RECONCILED = "RECONCILED"
    MISMATCH = "MISMATCH"
    INSUFFICIENT_EVIDENCE = "INSUFFICIENT_EVIDENCE"


class SavingsEvidenceState(str, Enum):
    ELIGIBLE_CANDIDATE = "ELIGIBLE_CANDIDATE"
    INSUFFICIENT_EVIDENCE = "INSUFFICIENT_EVIDENCE"


@dataclass(frozen=True, slots=True)
class EntitlementObservation:
    observation_id: str
    concept: EntitlementConcept
    tenant_id: str | None
    region_id: str
    reference: str
    account_identifier: str | None
    product_or_tier: str | None
    assignment_state: str | None
    assigned_at: object | None
    usage_period: str | None
    department_context: str | None
    explicit_unused_evidence: bool
    lineage: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class QuantityReconciliation:
    scope: str
    purchased: int | None
    assigned: int | None
    unassigned: int | None
    detail_assigned: int
    detail_unassigned: int
    state: QuantityReconciliationState
    explanation: str


@dataclass(frozen=True, slots=True)
class EntitlementAnalysis:
    observations: tuple[EntitlementObservation, ...]
    reconciliations: tuple[QuantityReconciliation, ...]
    purchased_license_quantity: int
    assigned_license_quantity: int
    unassigned_license_quantity: int
    access_entitlement_quantity: int
    assignment_utilization_percent: Decimal | None
    activity_utilization_percent: Decimal | None
    savings_state: SavingsEvidenceState
    savings_amount: Decimal | None
    blocked_reasons: tuple[str, ...]
    fingerprint: str


def analyze_entitlement_evidence(*, decomposition, classifications, evidence_sets, content: bytes):
    """Normalize bounded evidence without publishing ownership, savings, or source facts."""
    if decomposition.container.container_type != "XLSX":
        return _empty(decomposition.container.context)
    context = decomposition.container.context
    by_classification = {item.region_id: item for item in classifications}
    sheet_names = {item.sheet_id: item.name for item in decomposition.sheets}
    workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    observations = []
    summaries = {}
    detail = {}
    try:
        for region in decomposition.regions:
            location = region.location
            if not isinstance(location, SpreadsheetLocation):
                continue
            sheet_name = sheet_names[location.sheet_id]
            worksheet = workbook[sheet_name]
            rows = tuple(
                tuple(
                    worksheet.cell(row, column).value
                    for column in range(location.start_column, location.end_column + 1)
                )
                for row in range(location.start_row, location.end_row + 1)
            )
            leading = _leading(by_classification[region.region_id])
            if leading in {
                SemanticRegionType.ACCESS_ENTITLEMENT,
                SemanticRegionType.LICENSE_ASSIGNMENT,
            }:
                extracted = _assignment_rows(
                    context, region, sheet_name, location.start_row, rows, leading
                )
                observations.extend(extracted)
                detail[sheet_name] = extracted
            elif leading is SemanticRegionType.USAGE_DATA:
                observations.extend(
                    _usage_rows(context, region, sheet_name, location.start_row, rows)
                )
            else:
                found = _summary_quantities(rows)
                if found:
                    summaries.setdefault(sheet_name, {}).update(found)
    finally:
        workbook.close()
    reconciliations = []
    for sheet_name, items in detail.items():
        license_items = [
            item for item in items if item.concept is EntitlementConcept.LICENSE_ASSIGNMENT
        ]
        if not license_items:
            continue
        summary = summaries.get(sheet_name, {})
        assigned = sum(item.assignment_state == "ASSIGNED" for item in license_items)
        unassigned = sum(item.assignment_state == "UNASSIGNED" for item in license_items)
        purchased = _integer(summary.get("purchased"))
        stated_assigned = _integer(summary.get("assigned"))
        stated_unassigned = _integer(summary.get("unassigned"))
        state = QuantityReconciliationState.INSUFFICIENT_EVIDENCE
        explanation = "purchased and assignment summary evidence are incomplete"
        if purchased is not None:
            expected_unassigned = purchased - assigned
            if (
                (stated_assigned is None or stated_assigned == assigned)
                and (stated_unassigned is None or stated_unassigned == unassigned)
                and expected_unassigned == unassigned
            ):
                state = QuantityReconciliationState.RECONCILED
                explanation = "detail assignments and purchased quantity reconcile"
            else:
                state = QuantityReconciliationState.MISMATCH
                explanation = "detail assignments conflict with stated quantities"
        reconciliations.append(
            QuantityReconciliation(
                sheet_name,
                purchased,
                stated_assigned,
                stated_unassigned,
                assigned,
                unassigned,
                state,
                explanation,
            )
        )
    purchased_total = sum(item.purchased or 0 for item in reconciliations)
    assigned_total = sum(item.detail_assigned for item in reconciliations)
    unassigned_total = sum(item.detail_unassigned for item in reconciliations)
    access_total = sum(
        item.concept is EntitlementConcept.ACCESS_ENTITLEMENT for item in observations
    )
    utilization = (
        (Decimal(assigned_total) / Decimal(purchased_total) * 100).quantize(Decimal("0.01"))
        if purchased_total
        and all(item.state is QuantityReconciliationState.RECONCILED for item in reconciliations)
        else None
    )
    blocked = (
        "governed license price is unavailable",
        "governed currency authority is unavailable",
        "optimization eligibility is not established",
    )
    fingerprint = structural_fingerprint(
        ENTITLEMENT_VERSION, _scope(context), tuple(observations), tuple(reconciliations)
    )
    return EntitlementAnalysis(
        tuple(observations),
        tuple(reconciliations),
        purchased_total,
        assigned_total,
        unassigned_total,
        access_total,
        utilization,
        None,
        SavingsEvidenceState.INSUFFICIENT_EVIDENCE,
        None,
        blocked,
        fingerprint,
    )


def assert_entitlement_scope(left: EntitlementAnalysis, right: EntitlementAnalysis):
    left_tenants = {item.tenant_id for item in left.observations}
    right_tenants = {item.tenant_id for item in right.observations}
    if left_tenants and right_tenants and left_tenants != right_tenants:
        raise PermissionError("cross-tenant entitlement reconciliation is forbidden")


def _assignment_rows(context, region, sheet_name, start_row, rows, leading):
    if not rows:
        return []
    headers = [_normalize(item) for item in rows[0]]
    account_index = _column(headers, ("email", "userid", "accountid", "assignedto", "user"))
    product_index = _column(headers, ("licensetype", "license", "tier", "plan", "subscription"))
    status_index = _column(headers, ("status", "assignmentstatus", "state"))
    date_index = _column(headers, ("assigneddate", "assignmentdate", "date"))
    department_index = _column(headers, ("department", "team", "businessunit"))
    output = []
    seen = set()
    for offset, row in enumerate(rows[1:], 1):
        account = _value(row, account_index)
        status = str(_value(row, status_index) or "").strip().casefold()
        if not account and not status:
            continue
        identity = str(account or "").strip().casefold()
        product = str(_value(row, product_index) or "").strip() or None
        stable_key = (identity, product, status)
        identity_is_stable = bool(identity and "unassigned" not in identity)
        if identity_is_stable and stable_key in seen:
            continue
        if identity_is_stable:
            seen.add(stable_key)
        is_access = leading is SemanticRegionType.ACCESS_ENTITLEMENT
        concept = (
            EntitlementConcept.ACCESS_ENTITLEMENT
            if is_access
            else EntitlementConcept.LICENSE_ASSIGNMENT
        )
        unassigned = "unassigned" in identity or status in {"unused", "unassigned"}
        assignment_state = (
            "UNASSIGNED"
            if unassigned
            else ("INACTIVE" if status in {"inactive", "disabled"} else "ASSIGNED")
        )
        reference = f"{sheet_name}!{start_row + offset}"
        fingerprint = structural_fingerprint(
            ENTITLEMENT_VERSION + ".observation",
            region.region_id,
            reference,
            concept,
            identity,
            product,
            assignment_state,
        )
        output.append(
            EntitlementObservation(
                "entitlement-" + fingerprint[:24],
                concept,
                context.tenant_id,
                region.region_id,
                reference,
                identity or None,
                product,
                assignment_state,
                _value(row, date_index),
                None,
                str(_value(row, department_index) or "").strip() or None,
                status == "unused",
                (region.lineage.source_reference, region.lineage.physical_locator),
            )
        )
    return output


def _summary_quantities(rows):
    output = {}
    for row in rows:
        values = [item for item in row if item is not None]
        if len(values) < 2:
            continue
        label = _normalize(values[0])
        if "license" in label and "purchased" in label:
            output["purchased"] = values[1]
        elif label == "assigned":
            output["assigned"] = values[1]
        elif label in {"unused", "unassigned"}:
            output["unassigned"] = values[1]
    return output


def _usage_rows(context, region, sheet_name, start_row, rows):
    if not rows:
        return []
    headers = [_normalize(item) for item in rows[0]]
    account_index = _column(headers, ("email", "userid", "accountid", "user", "resourceid"))
    activity_index = _column(
        headers, ("lastactivity", "lastlogin", "eventcount", "usage", "quantity")
    )
    period_index = _column(headers, ("usageperiod", "billingperiod", "period", "date"))
    output = []
    for offset, row in enumerate(rows[1:], 1):
        account = str(_value(row, account_index) or "").strip() or None
        activity = _value(row, activity_index)
        period = str(_value(row, period_index) or "").strip() or None
        if account is None and activity is None:
            continue
        reference = f"{sheet_name}!{start_row + offset}"
        fingerprint = structural_fingerprint(
            ENTITLEMENT_VERSION + ".usage", region.region_id, reference, account, activity, period
        )
        output.append(
            EntitlementObservation(
                "entitlement-" + fingerprint[:24],
                EntitlementConcept.USAGE_OBSERVATION,
                context.tenant_id,
                region.region_id,
                reference,
                account,
                None,
                None,
                activity,
                period,
                None,
                False,
                (region.lineage.source_reference, region.lineage.physical_locator),
            )
        )
    return output


def _leading(classification):
    return next(
        (
            item.semantic_type
            for item in classification.candidates
            if not item.semantic_type.value.startswith("GENERIC_")
        ),
        SemanticRegionType.UNKNOWN,
    )


def _normalize(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _column(headers, aliases):
    return next((index for index, value in enumerate(headers) if value in aliases), None)


def _value(row, index):
    return row[index] if index is not None and index < len(row) else None


def _integer(value):
    try:
        return int(value) if value is not None else None
    except (TypeError, ValueError):
        return None


def _empty(context):
    fingerprint = structural_fingerprint(ENTITLEMENT_VERSION, _scope(context), ())
    return EntitlementAnalysis(
        (),
        (),
        0,
        0,
        0,
        0,
        None,
        None,
        SavingsEvidenceState.INSUFFICIENT_EVIDENCE,
        None,
        ("no bounded entitlement evidence",),
        fingerprint,
    )


def _scope(context):
    return (
        context.organization_id,
        context.tenant_id,
        context.prospect_id,
        context.analysis_id,
    )
