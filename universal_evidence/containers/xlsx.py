"""Bounded, business-neutral XLSX structural decomposition."""

from __future__ import annotations

import hashlib
import io
import posixpath
import zipfile
from dataclasses import dataclass, field, fields
from datetime import datetime
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from universal_evidence.contracts import EvidenceAnalysisContext
from universal_evidence.documents import (
    EvidenceRegion,
    ResourceDiagnostics,
    SecurityDiagnostic,
    SecurityFinding,
    SpreadsheetLocation,
    StructuralLineage,
    StructuralState,
    StructuralType,
    UploadContainer,
)
from universal_evidence.profiling.common import detect_header, nonempty
from universal_evidence.profiling.models import PrimitiveType, ProfilerConfig
from universal_evidence.profiling.primitive_types import observe_primitive

XLSX_DECOMPOSITION_VERSION = "pue-011c.xlsx.v1"


@dataclass(frozen=True, slots=True)
class XlsxDecompositionConfig:
    max_file_size: int = 25 * 1024 * 1024
    max_expanded_size: int = 100 * 1024 * 1024
    max_zip_members: int = 2_000
    max_compression_ratio: float = 250.0
    max_sheets: int = 50
    max_cells_inspected: int = 2_000_000
    max_regions: int = 250
    max_formulas: int = 250_000
    max_merged_ranges: int = 100_000
    max_formula_length: int = 512
    max_profile_labels: int = 100

    def __post_init__(self) -> None:
        if any(getattr(self, item.name) <= 0 for item in fields(self)):
            raise ValueError("XLSX decomposition limits must be positive")


@dataclass(frozen=True, slots=True)
class NamedRangeObservation:
    name: str
    destinations: tuple[str, ...]
    external: bool = False


@dataclass(frozen=True, slots=True)
class FormulaObservation:
    cell_reference: str
    expression: str
    cached_value_available: bool


@dataclass(frozen=True, slots=True)
class RegionProfile:
    header_row: int | None
    header_confidence: float
    header_method: str
    column_count: int
    populated_cell_count: int
    primitive_distribution: tuple[tuple[str, int], ...]
    number_formats: tuple[str, ...]
    formulas: tuple[FormulaObservation, ...]
    merged_ranges: tuple[str, ...]
    bounded_labels: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class SheetObservation:
    sheet_id: str
    name: str
    ordinal: int
    visibility: str
    populated_bounds: tuple[int, int, int, int] | None
    nonempty_cell_count: int
    formula_count: int
    merged_ranges: tuple[str, ...]
    hidden_rows: tuple[int, ...]
    hidden_columns: tuple[int, ...]
    named_ranges: tuple[NamedRangeObservation, ...]
    warnings: tuple[str, ...]
    region_ids: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class WorkbookDecomposition:
    container: UploadContainer
    decomposition_version: str
    sheets: tuple[SheetObservation, ...]
    documents: tuple = ()
    regions: tuple[EvidenceRegion, ...] = ()
    region_profiles: tuple[tuple[str, RegionProfile], ...] = ()
    status: StructuralState = StructuralState.DISCOVERED
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class _Cell:
    row: int
    column: int
    value: Any = field(compare=False, repr=False)
    primitive: PrimitiveType
    formula: str | None
    cached_available: bool
    number_format: str


def decompose_xlsx(
    *,
    context: EvidenceAnalysisContext,
    content: bytes,
    admitted_at: datetime,
    source_reference: str,
    filename: str | None = None,
    config: XlsxDecompositionConfig | None = None,
) -> WorkbookDecomposition:
    """Observe XLSX structure without interpreting business meaning."""
    policy = config or XlsxDecompositionConfig()
    content_fingerprint = hashlib.sha256(content).hexdigest()
    security, expanded_size, zip_warning, quarantine = _inspect_container(content, policy)
    resources = ResourceDiagnostics(
        len(content),
        expanded_size=expanded_size,
        limit_exceeded=zip_warning is not None,
        partial_reason=zip_warning,
    )
    initial_state = StructuralState.QUARANTINED if quarantine else (
        StructuralState.PARTIAL if zip_warning else StructuralState.DISCOVERED
    )
    container = UploadContainer(
        context,
        content_fingerprint,
        "XLSX",
        len(content),
        admitted_at,
        source_reference,
        security_status=initial_state,
        security_diagnostics=security,
        resource_diagnostics=resources,
        warnings=((zip_warning,) if zip_warning else ()),
        filename=filename,
    )
    if quarantine or zip_warning:
        return WorkbookDecomposition(
            container,
            XLSX_DECOMPOSITION_VERSION,
            (),
            status=initial_state,
            warnings=container.warnings,
        )
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=False, data_only=False, keep_links=False
        )
        cached_workbook = load_workbook(
            io.BytesIO(content), read_only=False, data_only=True, keep_links=False
        )
    except (InvalidFileException, OSError, ValueError, KeyError, zipfile.BadZipFile):
        malformed = SecurityDiagnostic(
            SecurityFinding.MALFORMED, "workbook could not be structurally opened"
        )
        rejected = _replace_container_state(container, StructuralState.QUARANTINED, malformed)
        return WorkbookDecomposition(
            rejected,
            XLSX_DECOMPOSITION_VERSION,
            (),
            status=StructuralState.QUARANTINED,
            warnings=("workbook could not be structurally opened",),
        )

    try:
        return _decompose_workbooks(workbook, cached_workbook, container, policy)
    finally:
        workbook.close()
        cached_workbook.close()


def _decompose_workbooks(workbook, cached_workbook, container, policy):
    sheets: list[SheetObservation] = []
    regions: list[EvidenceRegion] = []
    profiles: list[tuple[str, RegionProfile]] = []
    warnings: list[str] = []
    inspected = formulas = merged_count = 0
    limited = len(workbook.worksheets) > policy.max_sheets
    named = _named_ranges(workbook)
    for ordinal, worksheet in enumerate(workbook.worksheets[: policy.max_sheets], start=1):
        cached_sheet = cached_workbook.worksheets[ordinal - 1]
        sheet_id = f"{container.container_id}:sheet:{ordinal}"
        hidden_rows = tuple(
            index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden
        )
        hidden_columns = tuple(
            index for index, dimension in worksheet.column_dimensions.items() if dimension.hidden
        )
        merged = tuple(str(item) for item in worksheet.merged_cells.ranges)
        merged_count += len(merged)
        cells: dict[tuple[int, int], _Cell] = {}
        sheet_limited = False
        for row in worksheet.iter_rows():
            for cell in row:
                inspected += 1
                if inspected > policy.max_cells_inspected:
                    sheet_limited = limited = True
                    break
                if not nonempty(cell.value):
                    continue
                formula = (
                    str(cell.value)[: policy.max_formula_length]
                    if cell.data_type == "f"
                    else None
                )
                if formula:
                    formulas += 1
                cached = cached_sheet.cell(cell.row, cell.column).value
                cells[(cell.row, cell.column)] = _Cell(
                    cell.row,
                    cell.column,
                    cell.value,
                    observe_primitive(cell.value, formula=bool(formula)),
                    formula,
                    formula is not None and cached is not None,
                    str(cell.number_format or "General")[:128],
                )
            if sheet_limited:
                break
        if formulas > policy.max_formulas or merged_count > policy.max_merged_ranges:
            limited = sheet_limited = True
        components = _components(cells)
        for bounds in components:
            if len(regions) >= policy.max_regions:
                limited = sheet_limited = True
                break
            region, profile = _build_region(
                container, worksheet, sheet_id, cells, bounds, merged, len(regions)
            )
            regions.append(region)
            profiles.append((region.region_id, profile))
        sheet_regions = tuple(
            item.region_id
            for item in regions
            if isinstance(item.location, SpreadsheetLocation)
            and item.location.sheet_id == sheet_id
        )
        sheet_warnings = []
        if sheet_limited:
            sheet_warnings.append("resource limit reached during sheet decomposition")
        if worksheet.sheet_state != "visible" or hidden_rows or hidden_columns:
            sheet_warnings.append("hidden content observed")
        sheets.append(
            SheetObservation(
                sheet_id,
                worksheet.title,
                ordinal,
                worksheet.sheet_state,
                _bounds(cells),
                len(cells),
                sum(cell.formula is not None for cell in cells.values()),
                merged,
                hidden_rows,
                hidden_columns,
                tuple(item for item in named if _range_applies(item, worksheet.title)),
                tuple(sheet_warnings),
                sheet_regions,
            )
        )
    state = StructuralState.PARTIAL if limited else StructuralState.DISCOVERED
    if limited:
        warnings.append("workbook decomposition resource limit reached")
    hidden = any(
        sheet.visibility != "visible" or sheet.hidden_rows or sheet.hidden_columns
        for sheet in sheets
    )
    diagnostics = container.security_diagnostics
    if hidden:
        diagnostics += (
            SecurityDiagnostic(
                SecurityFinding.HIDDEN_CONTENT_PRESENT,
                "hidden workbook content observed",
            ),
        )
    resources = ResourceDiagnostics(
        container.size_bytes,
        expanded_size=container.resource_diagnostics.expanded_size,
        sheet_or_page_count=len(sheets),
        cell_count=inspected,
        region_count=len(regions),
        table_count=sum(item.structural_type is StructuralType.TABULAR for item in regions),
        limit_exceeded=limited,
        partial_reason=warnings[0] if warnings else None,
    )
    updated = UploadContainer(
        container.context,
        container.content_fingerprint,
        container.container_type,
        container.size_bytes,
        container.admitted_at,
        container.source_reference,
        security_status=state,
        security_diagnostics=diagnostics,
        resource_diagnostics=resources,
        warnings=tuple(warnings),
        filename=container.filename,
    )
    return WorkbookDecomposition(
        updated,
        XLSX_DECOMPOSITION_VERSION,
        tuple(sheets),
        regions=tuple(regions),
        region_profiles=tuple(profiles),
        status=state,
        warnings=tuple(warnings),
    )


def _inspect_container(content, policy):
    diagnostics: list[SecurityDiagnostic] = []
    if len(content) > policy.max_file_size:
        return (), None, "file size limit reached", False
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            expanded = sum(item.file_size for item in members)
            if len(members) > policy.max_zip_members:
                return (), expanded, "ZIP member limit reached", False
            if expanded > policy.max_expanded_size:
                return (), expanded, "expanded workbook size limit reached", False
            for item in members:
                normalized = posixpath.normpath(item.filename.replace("\\", "/"))
                if normalized.startswith("../") or PurePosixPath(normalized).is_absolute():
                    finding = SecurityDiagnostic(
                        SecurityFinding.MALFORMED, "unsafe ZIP member path observed"
                    )
                    return (finding,), expanded, None, True
                ratio = item.file_size / max(item.compress_size, 1)
                if ratio > policy.max_compression_ratio:
                    finding = SecurityDiagnostic(
                        SecurityFinding.RESOURCE_LIMIT_EXCEEDED,
                        "abnormal ZIP compression ratio observed",
                    )
                    return (finding,), expanded, None, True
                lower = item.filename.casefold()
                if lower.endswith("vbaproject.bin"):
                    diagnostics.append(
                        SecurityDiagnostic(SecurityFinding.MACRO_PRESENT, "macro content observed")
                    )
                elif "/externallinks/" in lower:
                    diagnostics.append(
                        SecurityDiagnostic(
                            SecurityFinding.EXTERNAL_LINKS_PRESENT,
                            "external link metadata observed; retrieval disabled",
                        )
                    )
                elif "/embeddings/" in lower or "/activex/" in lower:
                    diagnostics.append(
                        SecurityDiagnostic(
                            SecurityFinding.EMBEDDED_OBJECTS_PRESENT,
                            "embedded object content observed",
                        )
                    )
            dangerous = any(
                item.finding
                in {SecurityFinding.MACRO_PRESENT, SecurityFinding.EMBEDDED_OBJECTS_PRESENT}
                for item in diagnostics
            )
            return tuple(diagnostics), expanded, None, dangerous
    except zipfile.BadZipFile:
        malformed = SecurityDiagnostic(SecurityFinding.MALFORMED, "invalid XLSX container")
        return (malformed,), None, None, True


def _components(cells):
    if not cells:
        return ()
    rows = sorted({row for row, _ in cells})
    row_bands = _bands(rows)
    output = []
    for start_row, end_row in row_bands:
        columns = sorted(
            {column for row, column in cells if start_row <= row <= end_row}
        )
        for start_column, end_column in _bands(columns):
            if any(
                start_row <= row <= end_row and start_column <= column <= end_column
                for row, column in cells
            ):
                output.append((start_row, end_row, start_column, end_column))
    return tuple(sorted(output))


def _bands(values):
    if not values:
        return ()
    output = []
    start = previous = values[0]
    for value in values[1:]:
        if value > previous + 1:
            output.append((start, previous))
            start = value
        previous = value
    output.append((start, previous))
    return tuple(output)


def _build_region(container, worksheet, sheet_id, cells, bounds, merged, ordinal):
    start_row, end_row, start_column, end_column = bounds
    rows = [
        [
            cells.get((row, column)).value if (row, column) in cells else None
            for column in range(start_column, end_column + 1)
        ]
        for row in range(start_row, end_row + 1)
    ]
    local_config = ProfilerConfig(header_scan_rows=max(1, min(len(rows), 25)))
    header, confidence, method = detect_header(rows, local_config)
    header = start_row + header - 1 if header else None
    selected = [
        cell
        for (row, column), cell in cells.items()
        if start_row <= row <= end_row and start_column <= column <= end_column
    ]
    structural_type, structural_confidence = _classify(rows, selected, header)
    formats = tuple(sorted({cell.number_format for cell in selected}))
    formulas = tuple(
        FormulaObservation(
            worksheet.cell(cell.row, cell.column).coordinate,
            cell.formula or "",
            cell.cached_available,
        )
        for cell in selected
        if cell.formula is not None
    )
    distribution: dict[str, int] = {}
    for cell in selected:
        distribution[cell.primitive.value] = distribution.get(cell.primitive.value, 0) + 1
    intersecting_merged = tuple(
        item for item in merged if _range_intersects(worksheet[item], bounds)
    )
    labels = _bounded_structural_labels(
        rows,
        structural_type,
        (header - start_row) if header is not None else None,
    )
    profile = RegionProfile(
        header,
        confidence,
        method,
        end_column - start_column + 1,
        len(selected),
        tuple(sorted(distribution.items())),
        formats,
        formulas,
        intersecting_merged,
        labels,
    )
    location = SpreadsheetLocation(
        sheet_id,
        start_row,
        end_row,
        start_column,
        end_column,
        f"{worksheet.title}!{worksheet.cell(start_row, start_column).coordinate}:"
        f"{worksheet.cell(end_row, end_column).coordinate}",
    )
    locator = location.range_locator or f"region:{ordinal}"
    lineage = StructuralLineage(container.source_reference, container.container_id, locator)
    warnings = []
    if structural_type is StructuralType.UNKNOWN:
        warnings.append("structural classification unresolved")
    if formulas and not any(item.cached_value_available for item in formulas):
        warnings.append("formula cache unavailable or unverified")
    region = EvidenceRegion(
        container.context,
        container.container_id,
        None,
        location,
        structural_type,
        locator,
        XLSX_DECOMPOSITION_VERSION,
        structural_confidence,
        lineage,
        tuple(labels),
        profile_reference=structural_fingerprint_for_profile(profile),
        warnings=tuple(warnings),
        status=(
            StructuralState.UNRESOLVED
            if structural_type is StructuralType.UNKNOWN
            else StructuralState.DISCOVERED
        ),
    )
    return region, profile


def structural_fingerprint_for_profile(profile):
    from universal_evidence.documents import structural_fingerprint

    return structural_fingerprint(XLSX_DECOMPOSITION_VERSION + ".profile", profile)


def _bounded_structural_labels(rows, structural_type, local_header):
    candidates = []
    if local_header is not None and 0 <= local_header < len(rows):
        candidates = rows[local_header]
    elif structural_type in {StructuralType.KEY_VALUE, StructuralType.SUMMARY}:
        candidates = [row[0] for row in rows if row]
    return tuple(
        str(value).strip()[:80]
        for value in candidates
        if isinstance(value, str) and value.strip()
    )[:100]


def _classify(rows, cells, header):
    height = len(rows)
    width = max((len(row) for row in rows), default=0)
    if not cells or height == 0 or width == 0:
        return StructuralType.UNKNOWN, 0.0
    counts = [sum(nonempty(value) for value in row) for row in rows]
    strings = sum(cell.primitive is PrimitiveType.STRING for cell in cells)
    numeric = sum(
        cell.primitive in {PrimitiveType.INTEGER, PrimitiveType.DECIMAL, PrimitiveType.FORMULA}
        for cell in cells
    )
    tabular_rows = sum(count >= 2 for count in counts)
    if height >= 2 and width >= 3 and header is not None and tabular_rows >= 2:
        density = sum(counts) / (height * width)
        return StructuralType.TABULAR, max(0.55, min(0.95, 0.55 + 0.4 * density))
    if height >= 2 and width == 2 and all(count >= 1 for count in counts):
        first_strings = sum(
            bool(isinstance(row[0], str) and row[0].strip()) for row in rows
        )
        if first_strings >= max(1, height - 1):
            if numeric >= max(2, height // 2):
                return StructuralType.SUMMARY, 0.7
            return StructuralType.KEY_VALUE, 0.7
    if height == 1 and width > 1 and strings:
        return StructuralType.TEXT, 0.6
    if height >= 2 and width >= 2 and strings >= numeric:
        return StructuralType.FORM, 0.5
    return StructuralType.UNKNOWN, 0.35


def _bounds(cells):
    if not cells:
        return None
    rows = [item[0] for item in cells]
    columns = [item[1] for item in cells]
    return min(rows), min(columns), max(rows), max(columns)


def _named_ranges(workbook):
    output = []
    for name, defined in workbook.defined_names.items():
        try:
            destinations = tuple(
                f"{sheet}!{coordinate}" for sheet, coordinate in defined.destinations
            )
            external = any("[" in destination or "]" in destination for destination in destinations)
        except (AttributeError, TypeError, ValueError):
            destinations = ()
            external = True
        output.append(NamedRangeObservation(str(name), destinations, external))
    return tuple(output)


def _range_applies(item, sheet_name):
    quoted = sheet_name.replace("'", "''")
    return any(
        destination.startswith(f"{sheet_name}!") or destination.startswith(f"'{quoted}'!")
        for destination in item.destinations
    )


def _range_intersects(cells, bounds):
    start_row, end_row, start_column, end_column = bounds
    return any(
        start_row <= cell.row <= end_row and start_column <= cell.column <= end_column
        for row in cells
        for cell in row
    )


def _replace_container_state(container, state, diagnostic):
    return UploadContainer(
        container.context,
        container.content_fingerprint,
        container.container_type,
        container.size_bytes,
        container.admitted_at,
        container.source_reference,
        security_status=state,
        security_diagnostics=container.security_diagnostics + (diagnostic,),
        resource_diagnostics=container.resource_diagnostics,
        warnings=container.warnings,
        filename=container.filename,
    )
