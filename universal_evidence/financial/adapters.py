"""Bounded readback adapter from structural/semantic evidence into PUE-011F."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

import fitz
from openpyxl import load_workbook

from universal_evidence.documents import PdfLocation, SpreadsheetLocation
from universal_evidence.financial.documents import (
    FinancialDocumentAnalysis,
    FinancialDocumentInput,
    FinancialLineEvidence,
    FinancialValueEvidence,
    normalize_financial_document,
)


@dataclass(frozen=True, slots=True)
class AdaptedFinancialDocument:
    evidence: FinancialDocumentInput
    analysis: FinancialDocumentAnalysis


def analyze_decomposition_financial(
    *, decomposition, classifications, evidence_sets, content: bytes
) -> tuple[AdaptedFinancialDocument, ...]:
    """Run candidate invoice sets into F without publishing or persisting facts."""
    by_region = {item.region_id: item for item in decomposition.regions}
    output = []
    for candidate in evidence_sets:
        semantic_set = getattr(candidate, "evidence_set", candidate)
        if semantic_set.candidate_semantic_type != "INVOICE_METADATA":
            continue
        regions = tuple(by_region[item] for item in semantic_set.region_ids)
        lines, values = (
            _xlsx_evidence(decomposition, regions, content)
            if decomposition.container.container_type == "XLSX"
            else _pdf_evidence(regions, content)
        )
        context = semantic_set.context
        evidence = FinancialDocumentInput(
            semantic_set.evidence_set_id,
            context.organization_id,
            context.tenant_id,
            context.prospect_id,
            context.analysis_id,
            semantic_set.semantic_confidence or 0.0,
            semantic_set.region_ids,
            tuple(lines),
            tuple(values),
        )
        output.append(AdaptedFinancialDocument(evidence, normalize_financial_document(evidence)))
    return tuple(output)


def _xlsx_evidence(decomposition, regions, content):
    formulas = load_workbook(io.BytesIO(content), data_only=False, read_only=True)
    cached = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    names = {item.sheet_id: item.name for item in decomposition.sheets}
    lines, values = [], []
    try:
        for region in regions:
            location = region.location
            if not isinstance(location, SpreadsheetLocation):
                continue
            ws = formulas[names[location.sheet_id]]
            cached_ws = cached[names[location.sheet_id]]
            matrix = [
                [
                    ws.cell(row, column).value
                    for column in range(location.start_column, location.end_column + 1)
                ]
                for row in range(location.start_row, location.end_row + 1)
            ]
            header_index = _header_index(matrix)
            if header_index is not None:
                headers = [_norm(value) for value in matrix[header_index]]
                indices = {
                    name: _column(headers, aliases) for name, aliases in _LINE_COLUMNS.items()
                }
                for offset, row in enumerate(matrix[header_index + 1 :], header_index + 1):
                    label = (
                        str(row[indices["description"]] or "")
                        if indices["description"] is not None
                        else ""
                    )
                    if _summary_concept(label) or any(_summary_concept(value) for value in row):
                        _append_summary_row(values, region, ws, cached_ws, location, offset, row)
                        continue
                    amount_index = indices["amount"]
                    if amount_index is None or row[amount_index] is None:
                        continue
                    excel_row = location.start_row + offset
                    cells = [
                        ws.cell(excel_row, location.start_column + index)
                        for index in indices.values()
                        if index is not None
                    ]
                    formula_cells = tuple(
                        cell.coordinate for cell in cells if _is_formula(cell.value)
                    )
                    cached_cells = tuple(
                        cell.coordinate
                        for cell in cells
                        if _is_formula(cell.value) and cached_ws[cell.coordinate].value is not None
                    )
                    lines.append(
                        FinancialLineEvidence(
                            region.region_id,
                            f"{ws.title}!{excel_row}",
                            label or None,
                            _cell_value(row, indices["quantity"]),
                            _cell_value(row, indices["unit_price"]),
                            cached_ws.cell(excel_row, location.start_column + amount_index).value
                            if _is_formula(row[amount_index])
                            else row[amount_index],
                            _currency_candidate(cells),
                            formula_cells,
                            cached_cells,
                        )
                    )
            else:
                for offset, row in enumerate(matrix):
                    if len(row) >= 2 and row[0] is not None:
                        _append_value(
                            values,
                            region,
                            ws,
                            cached_ws,
                            location.start_row + offset,
                            location.start_column + 1,
                            row[0],
                        )
    finally:
        formulas.close()
        cached.close()
    return lines, values


def _pdf_evidence(regions, content):
    document = fitz.open(stream=content, filetype="pdf")
    lines, values = [], []
    try:
        for region in regions:
            location = region.location
            if not isinstance(location, PdfLocation) or not location.bounding_box:
                continue
            page = document[location.page_number - 1]
            rect = fitz.Rect(location.bounding_box)
            if region.structural_type.value == "TABULAR":
                table = min(
                    page.find_tables().tables, key=lambda item: _bbox_distance(item.bbox, rect)
                )
                rows = table.extract()
                header_index = _header_index(rows)
                if header_index is not None:
                    headers = [_norm(value) for value in rows[header_index]]
                    indices = {
                        name: _column(headers, aliases) for name, aliases in _LINE_COLUMNS.items()
                    }
                    for ordinal, row in enumerate(rows[header_index + 1 :], 1):
                        amount = _cell_value(row, indices["amount"])
                        if amount is not None:
                            lines.append(
                                FinancialLineEvidence(
                                    region.region_id,
                                    f"page:{location.page_number}:table-row:{ordinal}",
                                    str(_cell_value(row, indices["description"]) or "") or None,
                                    _cell_value(row, indices["quantity"]),
                                    _cell_value(row, indices["unit_price"]),
                                    amount,
                                    _token_currency(row),
                                )
                            )
            else:
                text = page.get_textbox(rect)
                parts = [item.strip() for item in text.splitlines() if item.strip()]
                for index, part in enumerate(parts):
                    concept = _summary_concept(part)
                    if concept and index + 1 < len(parts):
                        values.append(
                            FinancialValueEvidence(
                                region.region_id,
                                region.raw_structure_reference,
                                part,
                                parts[index + 1],
                                _token_currency(parts),
                                "text",
                            )
                        )
                    invoice_id = re.fullmatch(r"#([A-Z0-9][A-Z0-9-]{4,})", part.upper())
                    if not invoice_id:
                        invoice_id = re.search(
                            r"(?i)invoice\s*(?:no\.?|number)\s*[:#]?\s*([A-Z0-9][A-Z0-9-]{4,})",
                            part,
                        )
                    if invoice_id:
                        values.append(
                            FinancialValueEvidence(
                                region.region_id,
                                region.raw_structure_reference,
                                "Invoice Number",
                                invoice_id.group(1),
                            )
                        )
                iso = _iso_currency(parts)
                if iso:
                    values.append(
                        FinancialValueEvidence(
                            region.region_id,
                            region.raw_structure_reference,
                            "Currency",
                            iso,
                            iso,
                            "text",
                        )
                    )
    finally:
        document.close()
    return lines, values


_LINE_COLUMNS = {
    "description": ("description", "item", "service", "product"),
    "quantity": ("qty", "quantity", "units"),
    "unit_price": ("unitprice", "price", "rate", "unitcost"),
    "amount": ("amount", "lineamount", "extended", "charge"),
}


def _norm(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").casefold())


def _column(headers, aliases):
    return next((i for i, value in enumerate(headers) if value in aliases), None)


def _cell_value(row, index):
    return row[index] if index is not None and index < len(row) else None


def _is_formula(value):
    return isinstance(value, str) and value.startswith("=")


def _header_index(rows):
    for index, row in enumerate(rows):
        labels = [_norm(value) for value in row]
        if (
            _column(labels, _LINE_COLUMNS["amount"]) is not None
            and _column(labels, _LINE_COLUMNS["unit_price"]) is not None
        ):
            return index
    return None


def _summary_concept(value):
    text = _norm(value)
    return next(
        (
            label
            for label in ("subtotal", "totaldue", "tax", "discount", "credit", "fee", "balance")
            if label in text
        ),
        None,
    )


def _token_currency(values):
    text = " ".join(str(item) for item in values)
    iso = _iso_currency((text,))
    return iso or ("$" if "$" in text else None)


def _iso_currency(values):
    allowed = {
        "AED",
        "AUD",
        "BRL",
        "CAD",
        "CHF",
        "CNY",
        "EUR",
        "GBP",
        "HKD",
        "INR",
        "JPY",
        "KRW",
        "MXN",
        "NZD",
        "SEK",
        "SGD",
        "USD",
        "ZAR",
    }
    found = [
        item
        for item in re.findall(
            r"(?<![A-Z])[A-Z]{3}(?![A-Z])", " ".join(str(item) for item in values)
        )
        if item in allowed
    ]
    return found[0] if len(set(found)) == 1 else None


def _currency_candidate(cells):
    formats = [str(cell.number_format) for cell in cells]
    return _token_currency(formats)


def _bbox_distance(left, right):
    return sum(abs(float(a) - float(b)) for a, b in zip(left, right))


def _append_summary_row(values, region, ws, cached_ws, location, offset, row):
    label_index = next((i for i, item in enumerate(row) if _summary_concept(item)), None)
    if label_index is None or label_index + 1 >= len(row):
        return
    excel_row = location.start_row + offset
    cell = ws.cell(excel_row, location.start_column + label_index + 1)
    cached_cell = cached_ws[cell.coordinate]
    value = cached_cell.value if _is_formula(cell.value) else cell.value
    values.append(
        FinancialValueEvidence(
            region.region_id,
            f"{ws.title}!{cell.coordinate}",
            str(row[label_index]),
            value,
            _token_currency((cell.number_format,)),
            "number_format",
            cell.number_format,
            cell.value if _is_formula(cell.value) else None,
            cached_cell.value is not None,
        )
    )


def _append_value(values, region, ws, cached_ws, row, column, label):
    cell = ws.cell(row, column)
    cached_cell = cached_ws[cell.coordinate]
    source_value = cached_cell.value if _is_formula(cell.value) else cell.value
    currency = (
        _iso_currency((str(source_value).upper(),))
        if "currency" in _norm(label) or "ccy" in _norm(label)
        else None
    ) or _token_currency((cell.number_format,))
    values.append(
        FinancialValueEvidence(
            region.region_id,
            f"{ws.title}!{cell.coordinate}",
            str(label),
            source_value,
            currency,
            "number_format",
            cell.number_format,
            cell.value if _is_formula(cell.value) else None,
            cached_cell.value is not None,
        )
    )
