"""Upload-first PUE pilot admission and bounded structural region discovery."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from time import perf_counter
from typing import Any

from openpyxl import load_workbook

from services.prospect_data_intake_service import scan_upload
from universal_evidence.activation.fingerprint import fingerprint
from universal_evidence.capability import CapabilityScope
from universal_evidence.contracts import EvidenceAnalysisContext, EvidenceSource
from universal_evidence.profiling import FileProfile, profile_evidence
from universal_evidence.profiling.common import nonempty, row_regions


@dataclass(frozen=True, slots=True)
class StructuralTableRegion:
    sheet_id: str
    sheet_name: str
    region_kind: str
    start_row: int
    end_row: int
    header_row: int | None
    detail_record_count: int | None
    original_headers: tuple[str, ...]
    fingerprint: str


@dataclass(frozen=True, slots=True)
class EvidencePilotAdmission:
    scope: CapabilityScope
    source_id: str
    file_id: str
    evidence_fingerprint: str
    created_at: datetime
    expires_at: datetime
    profile: FileProfile
    regions: tuple[StructuralTableRegion, ...]
    legacy_analysis_reference: str | None
    authority: str
    fingerprint: str
    original_filename: str
    source_content: bytes = field(repr=False)


def admit_uploaded_evidence(
    tenant,
    *,
    filename: str,
    content: bytes,
    legacy_analysis=None,
    now=None,
    operations=None,
    operation_context=None,
):
    """Admit authorized evidence independently of legacy schema normalization."""
    from universal_evidence.operations import GovernedEventType, Severity, observe

    started = perf_counter()
    try:
        admission = _admit_uploaded_evidence(
            tenant, filename=filename, content=content, legacy_analysis=legacy_analysis, now=now
        )
    except Exception as exc:
        observe(
            operations,
            GovernedEventType.EVIDENCE_REJECTED,
            operation_context,
            audit=False,
            severity=Severity.WARNING,
            outcome="REJECTED",
            duration_ms=(perf_counter() - started) * 1000,
            attributes={"error_class": type(exc).__name__},
        )
        raise
    detail_records = sum(
        item.detail_record_count or 0
        for item in admission.regions
        if item.region_kind == "PRIMARY_DETAIL"
    )
    field_count = sum(
        len(item.original_headers)
        for item in admission.regions
        if item.region_kind == "PRIMARY_DETAIL"
    )
    safe = {
        "detail_records": detail_records,
        "fields": field_count,
        "legacy_compatibility": "NOTICE" if legacy_analysis is not None else "NOT_APPLICABLE",
    }
    references = {"evidence": admission.evidence_fingerprint}
    elapsed = (perf_counter() - started) * 1000
    observe(
        operations,
        GovernedEventType.EVIDENCE_ADMITTED,
        operation_context,
        audit=False,
        duration_ms=elapsed,
        references=references,
        attributes=safe,
    )
    observe(
        operations,
        GovernedEventType.EVIDENCE_PROFILED,
        operation_context,
        audit=False,
        duration_ms=elapsed,
        references=references,
        attributes=safe,
    )
    return admission


def _admit_uploaded_evidence(tenant, *, filename, content, legacy_analysis=None, now=None):
    scan = scan_upload(filename, content)
    now = now or datetime.now(timezone.utc)
    evidence_fingerprint = scan["sha256"]
    analysis_id = "pue-upload-analysis-" + fingerprint(
        tenant.tenant_id, tenant.audit_id, evidence_fingerprint
    )[:24]
    source_id = "pue-upload-source-" + evidence_fingerprint[:24]
    context = EvidenceAnalysisContext(
        analysis_id,
        source_id,
        tenant.tenant_id,
        None,
        None,
    )
    source = EvidenceSource(
        context,
        "authorized:prospect-upload",
        now,
        datetime.fromisoformat(tenant.expires_at),
    )
    profile = profile_evidence(source=source, filename=filename, content=content)
    regions = discover_structural_regions(profile, filename=filename, content=content)
    scope = CapabilityScope(analysis_id, tenant.tenant_id, None, None)
    legacy_reference = (
        str(getattr(legacy_analysis, "audit_id", "") or "") or None
    )
    identity = fingerprint(
        scope,
        source_id,
        profile.evidence_file.file_id,
        evidence_fingerprint,
        profile.structural_fingerprint,
        regions,
        legacy_reference,
    )
    return EvidencePilotAdmission(
        scope,
        source_id,
        profile.evidence_file.file_id,
        evidence_fingerprint,
        now,
        datetime.fromisoformat(tenant.expires_at),
        profile,
        regions,
        legacy_reference,
        "SHADOW / NON-AUTHORITATIVE",
        identity,
        filename,
        bytes(content),
    )


def admitted_source_rows(admission: EvidencePilotAdmission):
    """Read the immutable admitted source without consulting another tenant or upload."""
    return _source_rows(admission.original_filename, admission.source_content)


def discover_structural_regions(profile, *, filename, content):
    rows_by_sheet = _source_rows(filename, content)
    discovered = []
    for sheet_profile, rows in zip(profile.sheets, rows_by_sheet, strict=True):
        header_row = sheet_profile.sheet.header_row_candidate
        regions = row_regions(rows)
        primary_region = next(
            (
                bounds
                for bounds in regions
                if header_row is not None and bounds[0] <= header_row <= bounds[1]
            ),
            None,
        )
        if primary_region is not None:
            end_row = _detail_end(rows, header_row, primary_region[1])
            headers = tuple(
                str(value).strip() if value is not None else ""
                for value in rows[header_row - 1]
            )
            discovered.append(
                _region(
                    sheet_profile.sheet.sheet_id,
                    sheet_profile.sheet.original_name,
                    "PRIMARY_DETAIL",
                    header_row,
                    end_row,
                    header_row,
                    max(0, end_row - header_row),
                    headers,
                )
            )
        for start, end in regions:
            if primary_region is not None and (start, end) == primary_region:
                continue
            if header_row is not None and end < header_row:
                kind = "LEADING_METADATA"
            else:
                kind = "SECONDARY_SUMMARY"
            discovered.append(
                _region(
                    sheet_profile.sheet.sheet_id,
                    sheet_profile.sheet.original_name,
                    kind,
                    start,
                    end,
                    None,
                    None,
                    (),
                )
            )
    return tuple(discovered)


def _detail_end(rows: list[list[Any]], header_row: int, region_end: int):
    key_values = [
        rows[index - 1][0] if rows[index - 1] else None
        for index in range(header_row + 1, min(region_end, header_row + 20) + 1)
    ]
    populated = [value for value in key_values if nonempty(value)]
    numeric_key = (
        populated
        and sum(_is_number(value) for value in populated) / len(populated) >= 0.8
    )
    if numeric_key:
        for row_number in range(header_row + 1, region_end + 1):
            row = rows[row_number - 1]
            key = row[0] if row else None
            if any(nonempty(value) for value in row) and not _is_number(key):
                return row_number - 1
    return region_end


def _source_rows(filename, content):
    if Path(filename).suffix.lower() == ".csv":
        text = content.decode("utf-8-sig", errors="replace")
        return [[list(row) for row in csv.reader(io.StringIO(text))]]
    workbook = load_workbook(io.BytesIO(content), read_only=False, data_only=False)
    try:
        return [
            [list(row) for row in worksheet.iter_rows(values_only=True)]
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _region(sheet_id, sheet_name, kind, start, end, header, count, headers):
    identity = fingerprint(sheet_id, kind, start, end, header, count, headers)
    return StructuralTableRegion(
        sheet_id, sheet_name, kind, start, end, header, count, headers, identity
    )


def _is_number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)
